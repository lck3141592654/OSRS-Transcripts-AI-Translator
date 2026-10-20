#!/usr/bin/env python3
import sys
import json
import asyncio
from pathlib import Path
from datetime import datetime
import pandas as pd

sys.path.insert(0, str(Path(__file__).parent))

from glossary import load_glossary
from tm_matcher import match_and_fill
from llm_translator import prepare_sheet_translation, translate_sheet_phase, is_missing_translation
from llm_translator import save_session, delete_checkpoint_files, load_backup, get_debug_message_dir
from enforcer import enforce_async
from shared_pool import SharedBatchPool


SEP = "=" * 60
AUTO_GLOSSARY_SENTINEL = Path("__AUTO__")
WORKPLACE_DIR_NAME = "workplace"


def _get_workplace() -> Path:
    """取得腳本所在目錄同級的 workplace/ 目錄，不存在則自動建立。"""
    wp = Path(__file__).parent / WORKPLACE_DIR_NAME
    wp.mkdir(parents=True, exist_ok=True)
    return wp


def timestamp(msg: str):
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    print(f"  [{now}] {msg}")


def elapsed(start: datetime, label: str):
    secs = (datetime.now() - start).total_seconds()
    print(f"  [{label}] 耗時: {secs:.1f} 秒")

def filter_template_filled(pending, df):
    """模板已處理（_status == 已處理）的條目從待翻譯清單移除，避免重複送 LLM。"""
    if "_status" not in df.columns:
        return pending
    filled = set(df.index[df["_status"] == "已處理"])
    return [p for p in pending if p["_idx"] not in filled]


def validate_columns(frames: dict, sheet_names, required=("english", "translation")) -> bool:
    """檢查目標工作表是否具備必要欄位，避免後續 KeyError。

    frames：已載入的 {工作表: DataFrame}；呼叫端必須先讀好，這裡不再重複讀檔。
    """
    ok = True
    for sn in sheet_names:
        df = frames.get(sn)
        if df is None:
            print(f"  錯誤：找不到工作表「{sn}」的資料")
            ok = False
            continue
        missing = [c for c in required if c not in df.columns]
        if missing:
            print(f"  錯誤：工作表「{sn}」缺少必要欄位: {', '.join(missing)}")
            ok = False
    return ok


def get_frame(frames: dict | None, excel_path, sheet_name: str) -> pd.DataFrame:
    """取得工作表 DataFrame：優先用呼叫端已載入的快取（回傳副本），否則讀檔。

    回傳副本是必要的：呼叫端（_run_translation_pipeline、run_proofread）
    會用 df.update() 就地改寫，直接回傳快取物件會污染同一張表的其他階段。
    """
    if frames is not None and sheet_name in frames:
        return frames[sheet_name].copy()
    return pd.read_excel(excel_path, sheet_name=sheet_name, dtype=str)


def list_excel_files(directory: Path) -> list[Path]:
    if not directory.exists():
        return []
    return sorted([f for f in directory.glob("*.xlsx") if not f.name.startswith("~$")])


def choose_int(prompt: str, max_val: int, allow_zero: bool = False) -> int | None:
    while True:
        try:
            choice = input(prompt).strip()
            if not choice:
                return None
            idx = int(choice)
            if idx == 0 and allow_zero:
                return 0
            if 1 <= idx <= max_val:
                return idx
        except ValueError:
            pass
        print("  無效輸入，請重新輸入。")


def choose_from_list(items: list[str], title: str, allow_zero: bool = False) -> int | None:
    print(f"\n{SEP}")
    print(f"  {title}")
    print(SEP)
    for i, item in enumerate(items, 1):
        print(f"  [{i}] {item}")
    if allow_zero:
        print("  [0] 不使用 / 跳過")
    return choose_int("\n請輸入編號: ", len(items), allow_zero=allow_zero)


def choose_multi(items: list[str], title: str) -> list[str]:
    print(f"\n{SEP}")
    print(f"  {title}")
    print(SEP)
    for i, item in enumerate(items, 1):
        print(f"  [{i}] {item}")
    print("  [0] 全部選取")
    while True:
        try:
            choice = input("\n請輸入編號（可多個以逗號分隔）: ").strip()
            if choice == "0":
                return list(items)
            raw_indices = [int(x.strip()) for x in choice.split(",") if x.strip()]
            indices = list(dict.fromkeys(raw_indices))
            selected = [items[i-1] for i in indices if 1 <= i <= len(items)]
            if selected:
                return selected
        except (ValueError, IndexError):
            pass
        print("  無效輸入，請重新輸入。")


def step1() -> Path:
    """step1：從 workplace/ 選擇翻譯目標 Excel，並輸入語言代碼。"""
    workplace = _get_workplace()
    xlsx_files = list_excel_files(workplace)
    if not xlsx_files:
        print(f"\n錯誤：{workplace}/ 下沒有 Excel 檔案")
        print(f"請將翻譯目標 Excel 放入 {workplace}/ 後再執行。")
        input("\n按 Enter 關閉...")
        sys.exit(1)
    fnames = [f.name for f in xlsx_files]
    c = choose_from_list(fnames, f"請選擇目標 Excel 檔案（{workplace}/）：")
    if c is None:
        print("已取消。")
        input("\n按 Enter 關閉...")
        sys.exit(0)
    return xlsx_files[c - 1]


def step2(exclude_path: Path | None = None) -> tuple[Path | None, list[str] | None]:
    """step2：從 workplace/ 選擇術語庫 Excel。"""
    workplace = _get_workplace()
    xlsx_files = [f for f in list_excel_files(workplace)
                  if exclude_path is None or f.resolve() != exclude_path.resolve()]
    if not xlsx_files:
        inp = input(f"\n{workplace}/ 下沒有 Excel 檔案，輸入其他路徑或直接 Enter 跳過: ").strip()
        if inp:
            p = Path(inp)
            if p.exists():
                return p, None
        return None, None
    fnames = [f.name for f in xlsx_files]
    fnames.append("✨ 自動從目標 Excel 萃取（name/manual 工作表）")
    c = choose_from_list(fnames, "選擇術語庫 Excel（translation 有值的行即為術語）:", allow_zero=True)
    if c is None or c == 0:
        return None, None
    if c == len(fnames):
        return AUTO_GLOSSARY_SENTINEL, None
    gp = xlsx_files[c - 1]
    with pd.ExcelFile(gp) as xls:
        sheets = xls.sheet_names
    if len(sheets) == 1:
        print(f"  自動使用工作表: {sheets[0]}")
        return gp, sheets
    selected = choose_multi(sheets, "選擇哪些工作表作為術語來源：")
    return gp, selected


def step3(excel_path: Path) -> list[str]:
    """step3：選擇翻譯目標工作表。"""
    with pd.ExcelFile(excel_path) as xls:
        sheets = xls.sheet_names
    if len(sheets) == 1:
        return sheets
    return choose_multi(sheets, "選擇翻譯目標工作表：")

def step4_choose_sheet_mode(untranslated_count: int, total_rows: int) -> dict:
    """
    選擇單一工作表的翻譯範圍模式。
    回傳模式 dict:
      {"type": "all"}
      {"type": "first_n", "n": N}
      {"type": "range", "start": S, "end": E}
    """
    print(f"  未翻譯: {untranslated_count} 條")
    print(f"  [1] 全部未翻譯條目")
    print(f"  [2] 僅測試前 N 條")
    print(f"  [3] 指定行數範圍")
    while True:
        c = input("\n請選擇: ").strip()
        if c == "1":
            print(f"  選取全部 {untranslated_count} 條未翻譯")
            return {"type": "all"}
        elif c == "2":
            try:
                n = int(input("  請輸入 N: ").strip())
            except ValueError:
                print("  無效數字")
                continue
            if n < 1:
                print("  N 必須是 1 以上的正整數")
                continue
            print(f"  選取前 {n} 條未翻譯")
            return {"type": "first_n", "n": n}
        elif c == "3":
            try:
                rng = input("  請輸入行數範圍（如 100-500）: ").strip()
                parts = rng.split("-")
                if len(parts) != 2:
                    raise ValueError
                s = int(parts[0]) - 1
                e = int(parts[1])
                if not (0 <= s < e <= total_rows):
                    raise ValueError
                print(f"  選取行 {s + 1} 到 {e}")
                return {"type": "range", "start": s, "end": e}
            except (ValueError, IndexError):
                print(f"  無效範圍（請輸入 1 到 {total_rows} 之間的正向範圍）")
        else:
            print("  無效輸入")

    # 防呆：理論上不可到達（所有有效輸入皆在迴圈內 return）
    raise AssertionError("unreachable")


def step4_apply_mode(df_full: pd.DataFrame, mode: dict) -> pd.DataFrame:
    """根據預選模式對特定工作表的 DataFrame 套用範圍。"""
    untranslated_mask = df_full["translation"].apply(is_missing_translation)
    if mode["type"] == "all":
        return df_full[untranslated_mask].copy()
    elif mode["type"] == "first_n":
        n = mode["n"] if mode["n"] >= 0 else 0
        idx = list(df_full.index[untranslated_mask][:n])
        return df_full.loc[idx].copy() if idx else df_full.iloc[:0].copy()
    elif mode["type"] == "range":
        return df_full.iloc[mode["start"]:mode["end"]].copy()
    return df_full[untranslated_mask].copy()


def confirm(excel_path: Path, glossary_path: Path | None,
            sheets: list[str], total_rows: int) -> bool:
    """顯示翻譯摘要並要求確認。"""
    print(f"\n{SEP}")
    print("  即將執行以下操作：")
    print(SEP)
    print(f"  目標: {excel_path}")
    if glossary_path == AUTO_GLOSSARY_SENTINEL:
        print(f"  術語庫: 自動萃取（name/manual 工作表）")
    else:
        print(f"  術語庫: {glossary_path if glossary_path else '無（僅內建 ADD_LIST）'}")
    print(f"  工作表: {', '.join(sheets)}")
    print(f"  翻譯條數: {total_rows}")
    print(SEP)
    return input("\n確認執行？(Y/n): ").strip().lower() != "n"


def find_existing_progress() -> dict | None:
    """掃描 workplace/_checkpoint/session.json，偵測未完成的翻譯進度。"""
    workplace = _get_workplace()
    sf = workplace / "_checkpoint" / "session.json"
    if sf.exists():
        try:
            session = json.loads(sf.read_text(encoding="utf-8"))
            session["_workplace"] = str(workplace)
            return session
        except Exception:
            pass
    return None

def prompt_resume(session: dict) -> bool:
    """顯示續傳提示（支援新舊格式）。"""
    print(f"\n{SEP}")
    print("  偵測到未完成的翻譯進度！")
    print(SEP)
    excel_name = Path(session.get("excel_path", "?")).name
    print(f"  Excel: {excel_name}")

    if "current_sheet" in session:
        # 新版多工作表格式
        cs = session["current_sheet"]
        done = len(session.get("completed_sheets", []))
        total = len(session.get("pending_sheets", [])) + done
        print(f"  工作進度: {done}/{total} 工作表")
        print(f"  當前工作表: {cs}")
    else:
        # 舊版單工作表格式
        sn = session.get("sheet_name", "?")
        total = session.get("total_selected", "?")
        print(f"  工作表: {sn}")
        print(f"  條數: {total}")

    glossary_path = session.get("glossary_path")
    if glossary_path == "__AUTO__":
        print(f"  術語庫: 自動萃取（name/manual 工作表）")
    else:
        print(f"  術語庫: {Path(glossary_path).name if glossary_path else '無'}")
    print(SEP)
    return input("\n繼續翻譯？(Y/n): ").strip().lower() != "n"

def _build_session_data(excel_path, glossary_path, glossary_sheets,
                        current_sheet, completed_sheets, pending_sheets, sheet_configs,
                        accumulated_review_rows=None):
    data = {
        "excel_path": str(excel_path),
        "glossary_path": str(glossary_path) if glossary_path else None,
        "glossary_sheets": glossary_sheets,
        "current_sheet": current_sheet,
        "completed_sheets": list(completed_sheets),
        "pending_sheets": list(pending_sheets),
        "sheet_configs": sheet_configs,
        "timestamp": datetime.now().isoformat(),
    }
    if accumulated_review_rows:
        # 扁平化所有 DataFrame 為 dict 列表
        rows = []
        for df in accumulated_review_rows:
            for _, r in df.iterrows():
                rows.append(r.to_dict())
        data["accumulated_review_rows"] = rows
    return data

_FILLS = {
    "[glossary]": "D9D9D9",
    "[placeholder]": "BDD7EE",
    "[untranslated]": "F4CCCC",
    "[space]": "E2EFDA",
}


def write_review_report(path: Path, review_df: pd.DataFrame) -> None:
    """寫出審查報告並依問題類型上色。

    顏色直接在 ExcelWriter 的 worksheet 上套用，不再 to_excel → load_workbook
    → 上色 → save 的二次讀寫（5 萬列報告來回約 7 秒）。
    """
    from openpyxl.styles import Color, PatternFill
    fills = {k: PatternFill(start_color=Color(rgb=v), end_color=Color(rgb=v), fill_type="solid")
             for k, v in _FILLS.items()}
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        review_df.to_excel(writer, index=False, sheet_name="Sheet1")
        ws = writer.sheets["Sheet1"]
        issue_col = None
        for ci, cell in enumerate(ws[1], start=1):
            if cell.value == "issue":
                issue_col = ci
                break
        if issue_col is None:
            return
        for row in ws.iter_rows(min_row=2, max_col=ws.max_column, max_row=ws.max_row):
            v = str(row[issue_col - 1].value) if row[issue_col - 1].value else ""
            for prefix, fill in fills.items():
                if v.startswith(prefix):
                    for cell in row:
                        cell.fill = fill
                    break

def _write_outputs(excel_path: Path, workplace: Path, all_translated_dfs: dict,
                   all_review_rows: list[pd.DataFrame], overall_start: datetime):
    out_path = workplace / f"{excel_path.stem}_translated_output.xlsx"
    timestamp(f"寫入輸出檔案 {out_path.name}...")
    with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
        for sheet_name, df_data in all_translated_dfs.items():
            df_data.to_excel(writer, sheet_name=sheet_name, index=False)
    print(f"  輸出完成: {out_path}")

    if all_review_rows:
        combined_review = pd.concat(all_review_rows, ignore_index=True)
        # 與其他三份報告同口徑去重（證明用報告則共用 proofreader 的實作，避免第二份實作）
        from proofreader import _dedup_report_rows
        combined_review = _dedup_report_rows(
            combined_review, ["sheet_name", "english", "issue"], "review_report")
        report_path = workplace / "review_report.xlsx"
        write_review_report(report_path, combined_review)
        print(f"  審查報告已寫入: {report_path} ({len(combined_review)} 條)")

    delete_checkpoint_files(str(workplace))
    print(f"  已清除進度檔案")

    total_secs = (datetime.now() - overall_start).total_seconds()
    print(f"\n{SEP}")
    timestamp("管線執行完成")
    print(f"  總耗時: {total_secs:.1f} 秒")
    print(SEP)
    input("\n按 Enter 關閉...")

async def _run_translation_pipeline(excel_path, workplace, glossary, glossary_path, glossary_sheets,
                                    sheet_configs, pending_sheets, completed_sheets,
                                    all_translated_dfs, all_review_rows, frames=None):
    # Phase-barrier parallel translation:
    # Phase A local template match -> Phase B LLM translate (all sheets) -> Phase C enforce (all sheets).
    pool = SharedBatchPool()
    await pool.start()
    active = [c for c in pool.api_configs if not c.is_permanently_disabled]
    print(f"  已載入 {len(active)} 個可用 API：")
    for cfg in active:
        print(f"    {cfg.api_id} {cfg.model_provider}/{cfg.model} ({cfg.api_type}, 並發={cfg.parallel_limit})")
    try:
        # Phase A: local prep per sheet (restore backup + template match)
        sheet_states = {}
        for sheet_name in pending_sheets:
            timestamp(f"讀取 {sheet_name}...")
            df_full = get_frame(frames, excel_path, sheet_name)
            print(f"    總條數: {len(df_full)}")
            sheet_cfg = sheet_configs.get(sheet_name, {})
            selected_indices = set(sheet_cfg.get("selected_indices", []))
            if not selected_indices:
                print(f"  {sheet_name} 沒有選取任何條目，跳過")
                all_translated_dfs[sheet_name] = df_full
                completed_sheets.append(sheet_name)
                continue
            mask = df_full.index.isin(selected_indices) & (
                df_full["translation"].apply(is_missing_translation)
            )
            df = df_full[mask].copy()
            remaining = len(df)
            if remaining == 0:
                print(f"  {sheet_name} 已全部翻譯完成，跳過")
                all_translated_dfs[sheet_name] = df_full
                completed_sheets.append(sheet_name)
                continue
            t2 = datetime.now()
            timestamp("模板參數化比對...")
            df, pending = prepare_sheet_translation(df, str(workplace), sheet_name)
            df = match_and_fill(df, glossary)
            matched = int((df["_status"] == "已處理").sum())
            elapsed(t2, "模板比對")
            print(f"    模板匹配完成: {matched} 條已處理")
            pending = filter_template_filled(pending, df)
            pending_count = int(df["translation"].apply(is_missing_translation).sum())
            print(f"    待翻譯: {pending_count} 條")
            sheet_states[sheet_name] = {"df_full": df_full, "df": df, "pending": pending}

        # Phase B: LLM translation (all sheets in parallel, shared pool)
        t3 = datetime.now()
        timestamp("LLM 批次翻譯（全部工作表並行）...")
        await asyncio.gather(*[
            translate_sheet_phase(st["df"], glossary, str(workplace), st["pending"], sheet_name, pool)
            for sheet_name, st in sheet_states.items()
        ])
        elapsed(t3, "LLM 翻譯")
        save_session(str(workplace), _build_session_data(
            excel_path, glossary_path, glossary_sheets,
            pending_sheets[0] if pending_sheets else "",
            completed_sheets, pending_sheets, sheet_configs,
            accumulated_review_rows=all_review_rows))

        # Phase C: enforce (all sheets in parallel, shared pool)
        t4 = datetime.now()
        timestamp("術語強制後處理（全部工作表並行）...")
        results = await asyncio.gather(*[
            enforce_async(st["df"], glossary, str(workplace),
                          report_name=f"review_report_{sheet_name}.xlsx",
                          write_report=False, sheet_name=sheet_name, shared_pool=pool)
            for sheet_name, st in sheet_states.items()
        ], return_exceptions=True)
        for sheet_name, res in zip(sheet_states, results):
            if isinstance(res, BaseException):
                if isinstance(res, RuntimeError) and "all APIs permanently disabled" in str(res):
                    print("\n  ⚠️所有 API 已永久停用（累計 429 限流達門檻；只有一個可用 API 時門檻會放寬），無法繼續")
                    print("  已儲存部分翻譯進度，請檢查 API Key 後重新執行即可續傳")
                raise res
            df, review_df = res
            if not review_df.empty:
                review_df.insert(0, "sheet_name", sheet_name)
                all_review_rows.append(review_df)
            sheet_states[sheet_name]["df_full"].update(df)
            all_translated_dfs[sheet_name] = sheet_states[sheet_name]["df_full"]
            completed_sheets.append(sheet_name)
            timestamp(f"✅ {sheet_name} 完成")
        elapsed(t4, "術語強制後處理")

        # All sheets done: update session
        pending_sheets.clear()
        save_session(str(workplace), _build_session_data(
            excel_path, glossary_path, glossary_sheets,
            "", completed_sheets, pending_sheets, sheet_configs,
            accumulated_review_rows=all_review_rows))
    finally:
        await pool.close()


async def main():
    print(SEP)
    print("      混合翻譯管線 v1.5")
    print(f"      工作目錄: {_get_workplace()}")
    print(SEP)

    # ── 檢查上次的 debug 訊息 ──
    debug_dir = get_debug_message_dir(create=False)
    if debug_dir.exists():
        debug_files = sorted(debug_dir.glob("debug_*.json"))
        if debug_files:
            print(f"\n{'=' * 60}")
            print(f"  ⚠️ 上次執行有低翻譯率記錄")
            print(f"{'=' * 60}")
            for f in debug_files:
                try:
                    data = json.loads(f.read_text(encoding="utf-8"))
                    print(f"  ─ [{data.get('api_id', '?')}] {data.get('model', '?')} ─")
                    print(f"    批次大小: {data.get('batch_size')}, "
                          f"成功: {data.get('success')}/{data.get('returned')}, "
                          f"翻譯率: {data.get('rate', 0):.1%}")
                    preview = data.get("response_preview", "")
                    if preview:
                        print(f"    回傳內容:\n{preview}")
                except Exception as e:
                    print(f"  無法讀取 {f.name}: {e}")
            print(f"{'=' * 60}\n")
        import shutil
        shutil.rmtree(debug_dir)
        print(f"  已清除 debug 訊息\n")

    # ── 續傳檢測 ──
    exist_session = find_existing_progress()
    if exist_session:
        if not prompt_resume(exist_session):
            workplace = _get_workplace()
            delete_checkpoint_files(str(workplace))
            print("\n  已清除舊進度，開始新的翻譯。\n")
        else:
            # 解析 session（相容新舊格式）
            excel_path = Path(exist_session["excel_path"])
            glossary_path_str = exist_session.get("glossary_path")
            glossary_path = Path(glossary_path_str) if glossary_path_str else None
            glossary_sheets = exist_session.get("glossary_sheets")
            workplace = _get_workplace()

            if "current_sheet" in exist_session:
                # 新版多工作表格式
                pending_sheets = list(exist_session.get("pending_sheets", []))
                completed_sheets = list(exist_session.get("completed_sheets", []))
                sheet_configs = exist_session.get("sheet_configs", {})
            else:
                # 舊版單工作表格式 → 轉換為新版
                sn = exist_session.get("sheet_name", "")
                sel = list(exist_session.get("selected_indices", []))
                pending_sheets = [sn] if sn else []
                completed_sheets = []
                sheet_configs = {sn: {"selected_indices": sel}} if sn else {}

            # 載入術語庫
            overall_start = datetime.now()
            t0 = datetime.now()
            timestamp("載入術語庫...")
            if glossary_path_str == "__AUTO__":
                from glossary import auto_extract_glossary
                glossary = auto_extract_glossary(excel_path, workplace)
            else:
                glossary = load_glossary(glossary_path, glossary_sheets if glossary_path else None)
            elapsed(t0, "載入術語庫")
            print(f"  術語庫載入完成: {len(glossary)} 條")

            all_translated_dfs = {}
            all_review_rows = []

            # ── 還原已完成工作表資料 ──
            for sn in completed_sheets:
                timestamp(f"從檢查點還原工作表: {sn}...")
                df_full = pd.read_excel(excel_path, sheet_name=sn, dtype=str)
                backup = load_backup(str(workplace), sn)
                if backup:
                    restore_count = 0
                    bad_entries = 0
                    for idx_str, trans in backup.items():
                        # 檢查點檔損壞或人為編輯時 key 可能不是數字、值可能不是字串；
                        # 單筆壞資料不可讓整張工作表的還原失敗
                        # （與 llm_translator.prepare_sheet_translation 同口徑）
                        try:
                            idx = int(idx_str)
                        except (TypeError, ValueError):
                            bad_entries += 1
                            continue
                        if not isinstance(trans, str) or not trans.strip():
                            bad_entries += 1
                            continue
                        if idx in df_full.index:
                            df_full.at[idx, "translation"] = trans
                            restore_count += 1
                    print(f"    還原 {restore_count} 條翻譯")
                    if bad_entries:
                        print(f"    警告：檢查點有 {bad_entries} 筆無效資料（非數字 key 或非字串譯文），已略過")
                all_translated_dfs[sn] = df_full

            # ── 還原累積審查資料 ──
            if "accumulated_review_rows" in exist_session:
                rows = exist_session["accumulated_review_rows"]
                if rows:
                    all_review_rows.append(pd.DataFrame(rows))
                    print(f"    還原 {len(rows)} 條審查記錄")

            await _run_translation_pipeline(excel_path, workplace, glossary, glossary_path, glossary_sheets,
                                                    sheet_configs, pending_sheets, completed_sheets,
                                                    all_translated_dfs, all_review_rows)

            _write_outputs(excel_path, workplace, all_translated_dfs, all_review_rows, overall_start)
            delete_checkpoint_files(str(workplace))
            return

    # ── 全新翻譯路徑 ──
    excel_path = step1()
    workplace = _get_workplace()
    glossary_path, glossary_sheets = step2(exclude_path=excel_path)
    sheet_names = step3(excel_path)

    overall_start = datetime.now()
    t0 = datetime.now()
    timestamp("載入術語庫...")
    if glossary_path == AUTO_GLOSSARY_SENTINEL:
        from glossary import auto_extract_glossary
        glossary = auto_extract_glossary(excel_path, workplace)
    else:
        glossary = load_glossary(glossary_path, glossary_sheets)
    elapsed(t0, "載入術語庫")
    print(f"  術語庫載入完成: {len(glossary)} 條")

    # 收集各工作表的未翻譯統計（同時快取 DataFrame，減少重複讀取）
    sheet_untranslated = []
    cached_dfs = {}
    for sn in sheet_names:
        df_temp = pd.read_excel(excel_path, sheet_name=sn, dtype=str)
        if "translation" in df_temp.columns:
            uc = int(df_temp["translation"].apply(is_missing_translation).sum())
        else:
            uc = len(df_temp)
        sheet_untranslated.append((sn, uc))
        cached_dfs[sn] = df_temp

    # ← 原本在 :598-600 的驗證搬來這裡，並改吃快取
    if not validate_columns(cached_dfs, sheet_names):
        input("\n按 Enter 關閉...")
        return

    # 逐工作表選擇翻譯範圍
    sheet_configs = {}
    for sn, uc in sheet_untranslated:
        print(f"\n{SEP}")
        print(f"  設定工作表：{sn}（{uc} 條未翻譯）")
        print(SEP)
        sn_mode = step4_choose_sheet_mode(uc, len(cached_dfs[sn]))
        df_filtered = step4_apply_mode(cached_dfs[sn], sn_mode)
        sheet_configs[sn] = {
            "mode": sn_mode,
            "selected_indices": list(df_filtered.index) if len(df_filtered) > 0 else [],
        }

    # 顯示摘要並確認
    total_selected = sum(len(cfg["selected_indices"]) for cfg in sheet_configs.values())
    if total_selected == 0:
        print("\n所有工作表都沒有符合條件的條目，無需處理")
        input("\n按 Enter 關閉...")
        return
    if not confirm(excel_path, glossary_path, sheet_names, total_selected):
        print("  已取消。")
        input("\n按 Enter 關閉...")
        return

    # 儲存初始 session（多工作表格式）
    pending_sheets = list(sheet_names)
    save_session(str(workplace), _build_session_data(
        excel_path, glossary_path, glossary_sheets,
        pending_sheets[0] if pending_sheets else "",
        [], pending_sheets, sheet_configs))

    # ── 逐一處理所有工作表 ──
    all_translated_dfs = {}
    all_review_rows = []
    completed_sheets = []


    await _run_translation_pipeline(excel_path, workplace, glossary, glossary_path, glossary_sheets,
                                            sheet_configs, pending_sheets, completed_sheets,
                                            all_translated_dfs, all_review_rows, frames=cached_dfs)

    _write_outputs(excel_path, workplace, all_translated_dfs, all_review_rows, overall_start)
    delete_checkpoint_files(str(workplace))


if __name__ == "__main__":
    try:
        asyncio.run(main())
    except RuntimeError as e:
        if "all APIs permanently disabled" in str(e):
            print("\n  ⚠️ 所有 API 已永久停用（累計 429 限流達門檻；只有一個可用 API 時門檻會放寬），無法繼續。")
            print("  已儲存部分翻譯進度，請檢查 API Key 後重新執行即可續傳。")
            input("\n按 Enter 關閉...")
        else:
            print(f"\n發生錯誤: {e}")
            import traceback
            traceback.print_exc()
            input("\n按 Enter 關閉...")
    except Exception as e:
        print(f"\n發生錯誤: {e}")
        import traceback
        traceback.print_exc()
        input("\n按 Enter 關閉...")
